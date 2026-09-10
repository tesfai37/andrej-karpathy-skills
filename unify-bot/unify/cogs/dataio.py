"""Drag-and-drop import, and export in whatever shape somebody needs."""
from __future__ import annotations

import collections
import csv
import io
import zipfile
from datetime import datetime, timezone

import discord
from discord import app_commands
from discord.ext import commands

from .. import embeds, importer, store, views
from ..checks import admin_only
from ..config import ROLES, ROLE_LABEL

TABLES = {
    "members": ("SELECT id, gamertag, discord_id, legacy_name, active, created_at "
                "FROM members"),
    "achievements": (
        "SELECT m.gamertag, ma.role, ma.achievement_key, a.name, ma.mark, ma.granted_at "
        "FROM member_achievements ma JOIN members m ON m.id = ma.member_id "
        "JOIN achievements a ON a.key = ma.achievement_key ORDER BY m.gamertag, ma.role"),
    "scores": ("SELECT m.gamertag, s.trial_key, s.score, s.recorded_at FROM scores s "
               "JOIN members m ON m.id = s.member_id ORDER BY m.gamertag"),
    "parses": ("SELECT m.gamertag, p.label, p.dps, p.recorded_at FROM parses p "
               "JOIN members m ON m.id = p.member_id ORDER BY m.gamertag"),
    "role_map": "SELECT discord_role_id, kind, value, label FROM role_map",
    "guides": "SELECT category, topic, body, updated_at FROM guides ORDER BY category, topic",
    "trials": "SELECT key, name, short, emoji, sort, scored FROM trials ORDER BY sort",
    "catalog": ("SELECT key, name, trial_key, kind, sort, requires FROM achievements "
                "ORDER BY trial_key, sort"),
    "audit_log": "SELECT id, ts, actor_name, action, summary, undone FROM audit_log",
}


class DataIO(commands.Cog):
    """Import & export"""

    def __init__(self, bot):
        self.bot = bot

    # ------------------------------------------------------------------ export
    async def build_export(self, kind: str) -> discord.File:
        stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d")
        if kind == "db":
            # fold the write-ahead log in first, or the copy misses recent changes
            await self.bot.db.run("PRAGMA wal_checkpoint(TRUNCATE)")
            return discord.File(self.bot.db.path, filename=f"unify-{stamp}.sqlite3")

        data: dict[str, tuple[list[str], list[tuple]]] = {}
        for name, sql in TABLES.items():
            rows = await self.bot.db.all(sql)
            headers = list(rows[0].keys()) if rows else []
            data[name] = (headers, [tuple(r) for r in rows])

        if kind == "xlsx":
            return await self.build_workbook(stamp, data)

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, (headers, rows) in data.items():
                text = io.StringIO()
                writer = csv.writer(text)
                writer.writerow(headers)
                writer.writerows(rows)
                zf.writestr(f"{name}.csv", text.getvalue())
        buf.seek(0)
        return discord.File(buf, filename=f"unify-{stamp}-csv.zip")

    @staticmethod
    def cell(value):
        """Spreadsheets store numbers as floats, and a Discord id is far past the
        53 bits that survives - 700000000000000001 comes back as 7e+17 and the
        link is gone. Anything that big is written as text."""
        if isinstance(value, int) and not isinstance(value, bool) and abs(value) > 2**53:
            return str(value)
        return value

    async def wide_sheets(self) -> dict[str, tuple[list[str], list[list]]]:
        """The one-column-per-achievement layout the guild already uses, and the
        one /import reads. Exporting in this shape is what makes a backup
        restorable instead of merely readable."""
        db = self.bot.db
        members = await db.all(
            "SELECT id, gamertag, discord_id FROM members ORDER BY gamertag")
        keys = [r[0] for r in await db.all(
            "SELECT a.key FROM achievements a JOIN trials t ON t.key = a.trial_key "
            "ORDER BY t.sort, a.sort")]
        trials = [r[0] for r in await db.all(
            "SELECT key FROM trials WHERE scored = 1 ORDER BY sort")]
        labels = [r[0] for r in await db.all(
            "SELECT DISTINCT label FROM parses ORDER BY label")]

        marks: dict[tuple[int, str], dict[str, str]] = {}
        for r in await db.all(
                "SELECT member_id, role, achievement_key, mark FROM member_achievements"):
            marks.setdefault((r[0], r[1]), {})[r[2]] = r[3]
        scores = {(r[0], r[1]): r[2] for r in await db.all(
            "SELECT member_id, trial_key, score FROM scores")}
        parses = {(r[0], r[1]): r[2] for r in await db.all(
            "SELECT member_id, label, dps FROM parses")}

        sheets: dict[str, tuple[list[str], list[list]]] = {}
        for role, title in (("tank", "Tank"), ("healer", "Healer"),
                            ("dps", "DPS"), ("account", "AWA")):
            rows = []
            for m in members:
                held = marks.get((m[0], role), {})
                rows.append([m[1], str(m[2] or "")] + [held.get(k, "") for k in keys])
            sheets[title] = (["GamerTag", "DiscordName"] + keys, rows)

        sheets["Score"] = (
            ["GamerTag", "DiscordName"] + trials,
            [[m[1], str(m[2] or "")] + [scores.get((m[0], t), "") for t in trials]
             for m in members])
        if labels:
            sheets["Parse"] = (
                ["GamerTag", "DiscordName"] + labels,
                [[m[1], str(m[2] or "")] + [parses.get((m[0], p), "") for p in labels]
                 for m in members])
        return sheets

    async def build_workbook(self, stamp: str, reference: dict) -> discord.File:
        from openpyxl import Workbook

        wb = Workbook(write_only=True)          # 2,700 members x 60 columns x 4 roles
        for name, (headers, rows) in (await self.wide_sheets()).items():
            ws = wb.create_sheet(name[:31])
            ws.append(headers)
            for row in rows:
                ws.append(row)
        for name, (headers, rows) in reference.items():
            # These three are already carried by the wide sheets above. Leaving
            # them in also let "parses" match the importer's sheet-name rule, so a
            # re-import turned its label/dps/recorded_at headers into parse labels.
            if name in ("achievements", "scores", "parses"):
                continue
            ws = wb.create_sheet(name[:31])
            ws.append(headers)
            for row in rows:
                ws.append([self.cell(v) for v in row])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return discord.File(buf, filename=f"unify-{stamp}.xlsx")

    @app_commands.command(description="Download the whole database")
    @app_commands.choices(fmt=[
        app_commands.Choice(name="Excel workbook (.xlsx)", value="xlsx"),
        app_commands.Choice(name="CSV files (.zip)", value="csv"),
        app_commands.Choice(name="Raw database (.sqlite3)", value="db"),
    ])
    @admin_only()
    async def export(self, interaction: discord.Interaction,
                     fmt: app_commands.Choice[str] | None = None):
        await interaction.response.defer(thinking=True, ephemeral=True)
        kind = fmt.value if fmt else "xlsx"
        file = await self.build_export(kind)
        try:
            await interaction.followup.send(
                embed=embeds.success(
                    f"Export ready — {datetime.now(timezone.utc):%d %b %Y %H:%M} UTC"),
                file=file, ephemeral=True)
        except discord.HTTPException as exc:
            if exc.status != 413:
                raise
            return await interaction.followup.send(embed=embeds.error(
                "The export is bigger than this server's upload limit. Try "
                "`/export fmt:CSV files` — it compresses much smaller."), ephemeral=True)
        await self.bot.audit(interaction.user, "export", kind)

    # ------------------------------------------------------------------ import
    @app_commands.command(name="import", description="Upload a spreadsheet and I'll read it")
    @app_commands.describe(
        file="Drag your .xlsx or .csv straight onto Discord",
        role="Only needed for a CSV that doesn't say which role it is",
    )
    @app_commands.choices(role=[app_commands.Choice(name=ROLE_LABEL[r], value=r)
                                for r in ROLES] +
                               [app_commands.Choice(name="Account Wide", value="account")])
    @admin_only()
    async def import_cmd(self, interaction: discord.Interaction, file: discord.Attachment,
                         role: app_commands.Choice[str] | None = None):
        await interaction.response.defer(thinking=True)
        try:
            sheets = importer.read_file(await file.read(), file.filename)
        except Exception as exc:
            return await interaction.followup.send(embed=embeds.error(str(exc)))

        keys = {r[0] for r in await self.bot.db.all("SELECT key FROM achievements")}
        trials = {r[0] for r in await self.bot.db.all("SELECT key FROM trials")}
        plans = [
            importer.plan_sheet(name, headers, rows, keys, trials, role.value if role else "")
            for name, (headers, rows) in sheets.items()
        ]
        usable = [p for p in plans if p.kind != "skip"]

        e = embeds.base(self.bot.brand, f"📥  {file.filename}")
        lines = []
        for p in plans:
            if p.kind == "skip":
                lines.append(f"⏭️ **{p.name}** — skipped ({p.reason})")
            elif p.kind == "achievements":
                lines.append(f"✅ **{p.name}** — {p.rows} people, {len(p.value_cols)} "
                             f"achievement columns → **{ROLE_LABEL[p.role]}**")
            elif p.kind == "guides":
                lines.append(f"✅ **{p.name}** — {p.rows} reference entries")
            else:
                lines.append(f"✅ **{p.name}** — {p.rows} people, "
                             f"{len(p.value_cols)} {p.kind[:-1]} columns")
        e.description = "\n".join(lines)
        ignored = sorted({c for p in usable for c in p.ignored})
        if ignored:
            e.add_field(name="Columns I'll ignore",
                        value=", ".join(ignored[:20]) + ("…" if len(ignored) > 20 else ""),
                        inline=False)
        if not usable:
            e.color = 0xEF4444
            e.add_field(name="Nothing to import",
                        value="I need a **GamerTag** column and columns named after "
                              "achievements. Run `/export` and pick **Excel workbook** — "
                              "that file is already in the shape I read back.",
                        inline=False)
            return await interaction.followup.send(embed=e)

        e.add_field(
            name="​",
            value="People who aren't on the roster yet will be added, and existing records "
                  "are updated rather than wiped. A bulk import is **not** covered by "
                  "`/undo` — take an `/export` first if you want a restore point.",
            inline=False)
        view = views.Confirm(interaction.user.id, "Import", discord.ButtonStyle.success)
        message = await interaction.followup.send(embed=e, view=view, wait=True)
        view.message = message
        await view.wait()
        if not view.result:
            return await message.edit(embed=embeds.warn("Cancelled — nothing was imported."),
                                      view=None)

        await message.edit(embed=embeds.info(self.bot.brand, "Importing…"), view=None)
        report = await self.apply(sheets, usable, interaction.user)
        await self.bot.audit(interaction.user, "import", f"{file.filename} — {report}")
        await message.edit(embed=embeds.success(f"Imported **{file.filename}**\n{report}"))

    async def apply(self, sheets, plans: list[importer.SheetPlan], actor) -> str:
        """Bulk path: the catalog is read once and rows are written straight in,
        so a 500-person workbook doesn't turn into 20,000 audit entries."""
        db = self.bot.db
        catalog = await store.catalog(db)

        def with_prerequisites(keys: list[str]) -> set[str]:
            out, stack = set(), list(keys)
            while stack:
                key = stack.pop()
                if key in out or key not in catalog:
                    continue
                out.add(key)
                stack.extend(catalog[key]["requires"].split())
            return out

        new_people = granted = scored = parsed = legacy = guides = 0
        unreadable: collections.Counter = collections.Counter()
        conflicts: list[str] = []
        statements: list[tuple[str, tuple]] = []
        for plan in plans:
            _, rows = sheets[plan.name]

            if plan.kind == "guides":
                cols = {v: k for k, v in plan.value_cols.items()}
                for row in rows:
                    def cell(field):
                        i = cols[field]
                        return str(row[i]).strip() if i < len(row) and row[i] else ""
                    category, topic, body = cell("category"), cell("topic"), cell("body")
                    if category and topic and body:
                        statements.append((
                            "INSERT INTO guides(category, topic, body) VALUES(?,?,?) "
                            "ON CONFLICT(category, topic) DO UPDATE SET body = excluded.body, "
                            "updated_at = datetime('now')", (category, topic, body)))
                        guides += 1
                continue

            for row in rows:
                gamertag = (str(row[plan.gamertag_col] or "").strip()
                            if plan.gamertag_col < len(row) else "")
                if not gamertag:
                    continue
                discord_id = None
                if 0 <= plan.discord_col < len(row):
                    raw = str(row[plan.discord_col] or "").strip().strip("<@!>")
                    if raw.isdigit() and 15 <= len(raw) <= 25:   # a real snowflake
                        discord_id = int(raw)
                result = await store.upsert_member(db, gamertag, discord_id)
                member = result.member
                new_people += int(result.created)
                if result.conflict:
                    conflicts.append(result.conflict)

                if plan.kind == "achievements":
                    marks = {}
                    for col, key in plan.value_cols.items():
                        if col >= len(row):
                            continue
                        mark = importer.cell_mark(row[col])
                        if mark:
                            marks[key] = mark
                        elif not importer.cell_is_blank(row[col]):
                            unreadable[str(row[col]).strip()] += 1
                    # A prerequisite that wasn't spelled out is a plain clear.
                    for key in with_prerequisites(list(marks)):
                        statements.append((
                            "INSERT OR IGNORE INTO member_achievements"
                            "(member_id, role, achievement_key, mark, granted_by) "
                            "VALUES(?,?,?,?,?)",
                            (member.id, plan.role, key, marks.get(key, "X"), actor.id)))
                        granted += 1
                        legacy += int(marks.get(key) == "L")
                elif plan.kind == "scores":
                    for col, trial_key in plan.value_cols.items():
                        value = importer.cell_number(row[col]) if col < len(row) else None
                        if value:
                            statements.append((
                                "INSERT INTO scores(member_id, trial_key, score, recorded_by) "
                                "VALUES(?,?,?,?) ON CONFLICT(member_id, trial_key) DO UPDATE SET "
                                "score=excluded.score, recorded_at=datetime('now')",
                                (member.id, trial_key, value, actor.id)))
                            scored += 1
                else:
                    for col, label in plan.value_cols.items():
                        value = importer.cell_number(row[col]) if col < len(row) else None
                        if value:
                            statements.append((
                                "INSERT INTO parses(member_id, label, dps, recorded_by) "
                                "VALUES(?,?,?,?) ON CONFLICT(member_id, label) DO UPDATE SET "
                                "dps=excluded.dps, recorded_at=datetime('now')",
                                (member.id, label, value, actor.id)))
                            parsed += 1

        await db.run_many(statements)
        bits = [f"👥 {new_people} new members", f"🏅 {granted} achievements"]
        if scored:
            bits.append(f"🏆 {scored} scores")
        if parsed:
            bits.append(f"⚔️ {parsed} parses")
        if legacy:
            label = await self.bot.setting("mark_label")
            bits.append(f"🅛 {legacy} {label.lower()}")
        if guides:
            bits.append(f"📚 {guides} guide entries")
        summary = " • ".join(bits)
        if unreadable:
            shown = ", ".join(f"`{v}` ×{n}" for v, n in unreadable.most_common(6))
            summary += (f"\n\n📝 Cells I couldn't read (left out, nothing was guessed): "
                        f"{shown}")
        if granted and await self.bot.setting("role_sync"):
            summary += "\n\n🔗 Run `/sync all` to hand out the Discord roles for these."
        if conflicts:
            unique = list(dict.fromkeys(conflicts))
            summary += (f"\n\n⚠️ {len(conflicts)} row(s) needed a decision and were left "
                        "alone rather than overwriting somebody:\n"
                        + "\n".join(f"• {c}" for c in unique[:8])
                        + (f"\n• …and {len(unique) - 8} more" if len(unique) > 8 else ""))
        await store.log_action(db, actor.id, str(actor), "import", summary.split("\n")[0])
        return summary


async def setup(bot):
    await bot.add_cog(DataIO(bot))
