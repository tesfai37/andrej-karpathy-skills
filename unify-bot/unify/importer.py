"""Reads a dropped .xlsx or .csv and works out what's in it, so an admin never
has to reshape a spreadsheet before uploading it.

It understands the layout the guild already uses: one sheet per role, a GamerTag
column, a DiscordName column, and one column per achievement marked with an X."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

ROLE_SHEETS = {
    "tank": "tank", "tanks": "tank",
    "healer": "healer", "healers": "healer", "heal": "healer",
    "dps": "dps", "damage": "dps",
    "awa": "account", "account": "account", "accountwide": "account",
}
TRUTHY = {"x", "yes", "y", "true", "1", "✓", "✅", "done"}


def normalize(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


@dataclass
class SheetPlan:
    name: str
    kind: str                       # achievements | scores | parses | skip
    reason: str = ""                # why it was skipped
    role: str = ""
    gamertag_col: int = -1
    discord_col: int = -1
    value_cols: dict[int, str] = field(default_factory=dict)   # column index -> key/label
    ignored: list[str] = field(default_factory=list)
    rows: int = 0


def read_file(data: bytes, filename: str) -> dict[str, tuple[list[str], list[list]]]:
    """-> {sheet name: (headers, rows)}"""
    lower = filename.lower()
    if lower.endswith(".csv") or lower.endswith(".tsv"):
        delimiter = "\t" if lower.endswith(".tsv") else ","
        text = data.decode("utf-8-sig", errors="replace")
        table = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        if not table:
            return {}
        stem = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0]
        return {stem: (table[0], table[1:])}

    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = {}
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if not rows:
                continue
            out[ws.title] = ([str(c) if c is not None else "" for c in rows[0]], rows[1:])
        wb.close()
        return out

    raise ValueError("I can read .csv and .xlsx files. Save it as one of those and try again.")


def plan_sheet(name: str, headers: list[str], rows: list[list],
               achievement_keys: set[str], trial_keys: set[str],
               forced_role: str = "") -> SheetPlan:
    norm = [normalize(h) for h in headers]
    plan = SheetPlan(name=name, kind="skip", rows=len(rows))

    for i, h in enumerate(norm):
        if plan.gamertag_col < 0 and ("gamertag" in h or h in {"tag", "name", "player", "member"}):
            plan.gamertag_col = i
        elif plan.discord_col < 0 and "discord" in h:
            plan.discord_col = i

    if plan.gamertag_col < 0:
        plan.reason = "no GamerTag column"
        return plan

    sheet_key = normalize(name)
    if sheet_key.startswith("score"):
        plan.kind = "scores"
        plan.value_cols = {i: h for i, h in enumerate(norm) if h in trial_keys}
        if not plan.value_cols:
            plan.kind, plan.reason = "skip", "no columns matching a trial"
        return plan

    if sheet_key.startswith("parse"):
        plan.kind = "parses"
        plan.value_cols = {
            i: headers[i].strip() for i in range(len(norm))
            if i not in (plan.gamertag_col, plan.discord_col) and headers[i].strip()
        }
        if not plan.value_cols:
            plan.kind, plan.reason = "skip", "no parse columns"
        return plan

    role = forced_role or ROLE_SHEETS.get(sheet_key, "")
    if not role:
        plan.reason = f"can't tell which role '{name}' is for"
        return plan

    plan.role = role
    plan.value_cols = {i: h for i, h in enumerate(norm) if h in achievement_keys}
    plan.ignored = [
        headers[i] for i, h in enumerate(norm)
        if h and i not in plan.value_cols and i not in (plan.gamertag_col, plan.discord_col)
    ]
    if not plan.value_cols:
        plan.kind, plan.reason = "skip", "no columns matching a known achievement"
    else:
        plan.kind = "achievements"
    return plan


def cell_is_true(value) -> bool:
    return str(value or "").strip().lower() in TRUTHY


NUMBER_RE = re.compile(r"([0-9]*\.?[0-9]+)\s*([km])?\s*(?:dps|score|pts?)?")


def cell_number(value) -> int | None:
    """Accepts 112000, "112,000", "112k", "1.2m", "112k dps" - the ways scores
    actually get typed - and refuses everything else.

    It deliberately does not go digit-hunting inside prose: a stray notes column
    reading "see notes 2024" must not import as a parse of 2024."""
    text = ("" if value is None else str(value)).strip().lower().replace(",", "")
    if not text:
        return None
    match = NUMBER_RE.fullmatch(text)
    if match is None:
        return None
    number = float(match.group(1)) * {"k": 1_000, "m": 1_000_000}.get(match.group(2), 1)
    return int(number)
